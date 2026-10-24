"""Genera el paquete de revision clinica para nutricionista o medico (plan de fiabilidad, Fases 4 y 7).

Produce un .xlsx con:
  - Instrucciones: que revisar, como leer los scores y que cuenta como error grave.
  - Casos: >=100 perfiles sinteticos estratificados, corridos end-to-end por /recommend (sin BD), con la salida
    del sistema y columnas de juicio con listas desplegables.
  - Una hoja por tabla de reglas (condicion, seguridad, interacciones, vinculos condicion-componente, etapa de
    vida, reglas de ingrediente) con decision aprobar/modificar/rechazar.
  - Firma: datos del profesional y hash de cada archivo revisado, para que la revision quede versionada.
y un manifest JSON con los mismos hashes. `importar_revision.py` lee el .xlsx completado.

Los casos son sinteticos: no contienen datos de personas reales.

    python scripts/clinical/generar_paquete_revision.py --cases 120 --out-dir data/reports/clinical_review
"""

from __future__ import annotations

import argparse
import dataclasses
import hashlib
import json
import random
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import pandas as pd

KNOWLEDGE_DIR = ROOT / "data" / "knowledge"
PACKAGE_VERSION = "clinical_review_v1"
DEFAULT_OUT_DIR = ROOT / "data" / "reports" / "clinical_review"

# Hojas de reglas: (nombre de hoja, CSV, columnas que identifican la fila de forma estable).
RULE_SHEETS: list[tuple[str, str, list[str]]] = [
    ("Reglas condicion", "condition_rules.csv", ["condition_code", "field", "operator", "value"]),
    ("Reglas seguridad", "safety_rules.csv", ["rule_code"]),
    ("Interacciones", "component_interaction_rules.csv", ["component_id", "interaction_type", "trigger_field", "trigger_value"]),
    ("Vinculos condicion-comp", "condition_component_links.csv", ["condition_code", "component_id"]),
    ("Etapa de vida", "component_life_stage_guidance.csv", ["component_id", "age_min", "age_max", "sex", "pregnancy", "lactation"]),
]
INGREDIENT_SHEET = "Reglas ingrediente"
REVIEWED_FILES = [
    "data/knowledge/condition_rules.csv",
    "data/knowledge/safety_rules.csv",
    "data/knowledge/component_interaction_rules.csv",
    "data/knowledge/condition_component_links.csv",
    "data/knowledge/component_life_stage_guidance.csv",
    "data/knowledge/conditions.csv",
    "data/knowledge/biomarkers.csv",
    "data/knowledge/sources.csv",
    "app/domains/catalog/seguridad_productos.py",
    "models/runtime/condition_mvp_model.pkl",
    "models/runtime/modelo2_artifacts.pkl",
]

CASE_REVIEW_COLUMNS: dict[str, list[str]] = {
    "rev_condiciones": ["de_acuerdo", "parcial", "en_desacuerdo"],
    "rev_condiciones_faltantes": [],
    "rev_condiciones_sobrantes": [],
    "rev_recomendaciones": ["de_acuerdo", "parcial", "en_desacuerdo"],
    "rev_componentes_inadecuados": [],
    "rev_componentes_faltantes": [],
    "rev_seguridad": ["correcta", "insuficiente", "excesiva"],
    "rev_error_grave": ["no", "si"],
    "rev_comentario": [],
}
RULE_REVIEW_COLUMNS: dict[str, list[str]] = {
    "rev_decision": ["aprobar", "modificar", "rechazar", "no_se"],
    "rev_cambio_propuesto": [],
    "rev_fuente_sugerida": [],
    "rev_comentario": [],
}
SIGNATURE_FIELDS = ["Nombre completo", "Profesion", "N. de colegiatura (CNP/CMP)", "Fecha de revision", "Firma (texto o escaneo adjunto)"]

INSTRUCTIONS = [
    "Paquete de revision clinica de SupleMatch",
    "",
    "Objetivo: validar que el sistema detecta riesgos nutricionales plausibles, recomienda componentes adecuados y "
    "bloquea o advierte cuando corresponde. Su revision reemplaza al benchmark interno de 27 casos.",
    "",
    "1) Hoja 'Casos' (prioridad). Cada fila es un perfil sintetico (no es una persona real) y lo que el sistema le "
    "mostraria. Complete las columnas rev_*: ",
    "   - rev_condiciones: acuerdo con las condiciones detectadas. Anote faltantes/sobrantes con el nombre.",
    "   - rev_recomendaciones: acuerdo con los componentes sugeridos (se muestran los 10 primeros; la lista completa "
    "esta en la hoja 'Recomendaciones'). Anote inadecuados/faltantes. Si la lista es demasiado larga para ser util, "
    "digalo en rev_comentario.",
    "   - rev_seguridad: 'insuficiente' si faltan advertencias o bloqueos; 'excesiva' si bloquea sin motivo.",
    "   - rev_error_grave = si cuando seguir la recomendacion podria causar dano (p. ej. interaccion con "
    "medicacion, dosis riesgosa, embarazo, enfermedad renal) o retrasar una consulta necesaria.",
    "",
    "2) Hojas de reglas. Una fila por regla del sistema. rev_decision: aprobar / modificar (describa el cambio en "
    "rev_cambio_propuesto) / rechazar / no_se. Si conoce una fuente mejor (guia, revision Cochrane, DOI), anotela "
    "en rev_fuente_sugerida.",
    "",
    "3) Hoja 'Firma'. Complete sus datos. La tabla de hashes identifica la version exacta revisada: si el equipo "
    "cambia una regla despues, esa regla debe volver a revisarse.",
    "",
    "Como leer los scores:",
    "   - indice de riesgo (0-1): indice basado en reglas y un modelo entrenado con datos sinteticos. NO es una "
    "probabilidad clinica.",
    "   - 'probabilidad calibrada (NHANES)': modelo con datos reales de EE. UU.; no validado en Peru.",
    "   - Nivel de seguridad: normal / caution (advertencias) / medical_review_required (compra directa bloqueada).",
    "",
    "Tiempo estimado: 2-3 min por caso, 30-45 min por hoja de reglas. Puede hacerlo en varias sesiones; guarde el "
    "archivo con el mismo nombre.",
]

SAFETY_CONDITIONS = [
    "embarazo_lactancia", "enfermedad_renal", "enfermedad_hepatica", "problema_tiroideo", "anticoagulantes",
    "medicacion_cronica", "diabetes_medicada", "inmunosupresion", "sedantes", "hipertension",
]
TOP_RECS_IN_CASE = 10
LAB_CODES = ["vitamin_d", "b12", "ferritin", "hemoglobin", "glucose", "tsh", "egfr", "magnesium", "folate", "ldl"]


def sha256_12(path: Path) -> str | None:
    if not path.exists():
        return None
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()[:12]


def file_manifest(root: Path = ROOT) -> list[dict[str, Any]]:
    return [{"file": rel, "sha256_12": sha256_12(root / rel)} for rel in REVIEWED_FILES]


def row_key(sheet: str, row: dict[str, Any], key_cols: list[str]) -> str:
    """Id estable de una regla para cruzar la revision con el CSV aunque cambie el orden de filas."""
    raw = sheet + "|" + "|".join(str(row.get(col, "")) for col in key_cols)
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()[:10]


# ---------------------------------------------------------------------------
# Casos
# ---------------------------------------------------------------------------

def _age_range(age: int) -> str:
    if age < 18:
        return "menos_18"
    if age <= 30:
        return "18_30"
    if age <= 50:
        return "31_50"
    return "mas_50"


def _lab_value(code: str, rng: random.Random) -> dict[str, Any] | None:
    from app.domains.labs.biomarcadores import BIOMARKERS

    definition = BIOMARKERS.get(code)
    if definition is None:
        return None
    low, high = definition.low, definition.high
    band = rng.choice(["low", "low", "normal", "high"])
    if band == "low" and low:
        value = low * rng.uniform(0.45, 0.9)
    elif band == "high" and high:
        value = high * rng.uniform(1.1, 1.6)
    else:
        lo = low or (high or 1) * 0.5
        hi = high or lo * 2
        value = rng.uniform(lo, hi)
    return {"code": code, "value": round(max(value, 0.1), 2), "unit": definition.unit}


# Valores de la escala semanal del frontend desde el contrato 2026-10-24.1 (SCALE_PRESETS.weekly_food):
# no consumo, 1/sem, 2-3/sem, 4-6/sem, 1 al dia, 2+ al dia. La escala anterior terminaba en "5+/sem" (=7) y nadie
# podia declarar dos porciones diarias de lacteos o proteina.
WEEKLY_SCALE = [0, 1, 3, 5, 7, 14]
FOOD_FIELDS = [
    "red_meat_servings_week", "poultry_servings_week", "fish_servings_week",
    "eggs_servings_week", "legume_servings_week", "dairy_servings_week",
]


def _food_block(rng: random.Random, tipo_dieta: str) -> dict[str, Any]:
    """Bloque de alimentos completo y coherente con el tipo de dieta (como lo llenaria la app).

    Antes solo se generaba pescado y el estimador de proteina trataba el resto como 0: todos los casos salian
    con proteina critica. Un bloque parcial no representa a nadie: o se responde entero o no se responde.
    """
    pick = lambda options: rng.choice(options)  # noqa: E731
    block = {
        "red_meat_servings_week": pick([0, 1, 3, 5]),
        "poultry_servings_week": pick([1, 3, 5, 7, 14]),
        "fish_servings_week": pick([0, 1, 1, 3, 5]),
        "eggs_servings_week": pick([1, 3, 5, 7]),
        "legume_servings_week": pick([0, 1, 3, 5]),
        "dairy_servings_week": pick([0, 1, 3, 7, 14]),
    }
    if tipo_dieta in {"vegetariano", "vegano", "pescetariano"}:
        block["red_meat_servings_week"] = block["poultry_servings_week"] = 0
        block["no_meat"] = True
    if tipo_dieta in {"vegetariano", "vegano"}:
        block["fish_servings_week"] = 0
        block["legume_servings_week"] = pick([3, 5, 7, 14])
    if tipo_dieta == "vegano":
        block["eggs_servings_week"] = block["dairy_servings_week"] = 0
    block["fruit_veg_servings_day"] = pick([0, 1, 2, 3, 5])
    return block


def _optional_context(rng: random.Random, payload: dict[str, Any]) -> dict[str, Any]:
    """Preguntas nuevas del contrato 2026-10-24.1; cada una queda sin responder ~30 % de las veces."""
    answered = lambda: rng.random() < 0.7  # noqa: E731
    extra: dict[str, Any] = {}
    if payload["fatiga"] in {"a_menudo", "siempre"} and answered():
        extra["fatiga_duracion"] = rng.choice(["menos_2_semanas", "2_semanas_3_meses", "mas_3_meses"])
        extra["fatiga_mejora_descanso"] = rng.choice(["si", "no", "a_veces"])
    if answered():
        extra["protector_solar"] = rng.choice(["nunca", "a_veces", "siempre"])
        extra["piel_expuesta"] = rng.choice(["cara_manos", "brazos", "brazos_piernas"])
    pregnant = "embarazo_lactancia" in payload["condiciones_seguridad"]
    if payload["sexo"] == "femenino" and 14 <= payload["age_years"] <= 50 and not pregnant and answered():
        extra["menstruacion_abundante"] = rng.choice(["no", "no", "si"])
    if payload["sexo"] == "femenino" and 18 <= payload["age_years"] <= 45 and answered():
        extra["parto_ultimo_anio"] = rng.choice(["no", "no", "no", "si"])
    if answered():
        extra["dona_sangre"] = rng.choice(["no", "no", "no", "1_al_anio", "2_o_mas_al_anio"])
    if payload["tipo_dieta"] in {"vegano", "vegetariano"} and answered():
        extra["anios_dieta_vegetal"] = rng.choice(["menos_1", "1_3", "mas_3"])
    if answered():
        extra["medicamentos_b12"] = rng.choice([[], [], [], ["metformina"], ["antiacidos_ibp"]])
        extra["alimentos_fortificados_b12"] = rng.choice(["no", "si", "no_se"])
    # Paso "Contexto adicional opcional" de la app.
    if answered():
        extra["screen_hours_day"] = rng.choice([2, 4, 6, 8, 10])
        extra["water_intake_l_day"] = rng.choice([0.8, 1.2, 1.5, 2.0, 2.5])
        extra["heavy_sweat_days_week"] = rng.choice([0, 0, 1, 3, 5])
        extra["digestive_discomfort"] = rng.choice(["nunca", "nunca", "leve", "moderado", "frecuente"])
    extra["analisis_ultimo_anio"] = "si_a_mano" if payload.get("lab_results") else rng.choice(["no", "no_se", "si_sin_resultados"])
    return extra


def build_case_payloads(n_cases: int, seed: int = 2026) -> list[dict[str, Any]]:
    """Perfiles sinteticos estratificados: cuotas fijas para seguridad, etapas de vida, dietas y labs; el resto al azar."""
    rng = random.Random(seed)
    cases: list[dict[str, Any]] = []

    # Cuotas: cada condicion de seguridad 4 veces, 4 menores, 8 mayores de 60, 15 veg, 25 con labs.
    forced: list[dict[str, Any]] = []
    for condition in SAFETY_CONDITIONS:
        for _ in range(4):
            forced.append({"condiciones_seguridad": [condition]})
    forced += [{"age_years": rng.randint(14, 17)} for _ in range(4)]
    forced += [{"age_years": rng.randint(61, 82)} for _ in range(8)]
    forced += [{"tipo_dieta": rng.choice(["vegano", "vegetariano"])} for _ in range(15)]
    forced += [{"_labs": rng.sample(LAB_CODES, k=rng.randint(2, 4))} for _ in range(25)]

    for index in range(n_cases):
        override = forced[index] if index < len(forced) else {}
        sexo = "femenino" if "embarazo_lactancia" in override.get("condiciones_seguridad", []) else rng.choice(["femenino", "masculino"])
        age = override.get("age_years") or rng.randint(19, 60)
        height = rng.randint(150, 175) if sexo == "femenino" else rng.randint(160, 188)
        weight = round(rng.uniform(48, 95), 1)
        payload: dict[str, Any] = {
            "edad_rango": _age_range(age),
            "age_years": age,
            "sexo": sexo,
            "weight_value": weight, "weight_unit": "kg",
            "height_value": height, "height_unit": "cm",
            "horas_sueno": rng.choice(["menos_5h", "5_7h", "5_7h", "7_9h", "7_9h", "mas_9h"]),
            "frecuencia_ejercicio": rng.choice(["casi_nunca", "1_2_semana", "3_4_semana", "diario"]),
            "dieta": rng.choice(["poco_variada", "regular", "bastante_variada", "muy_balanceada"]),
            "fatiga": rng.choice(["siempre", "a_menudo", "a_veces", "casi_nunca"]),
            "exposicion_solar": rng.choice(["menos_15min", "15_30min", "30_60min", "mas_1h"]),
            "frecuencia_enfermedad": rng.choice(["muy_seguido", "3_4_anio", "1_2_anio", "casi_nunca"]),
            "estres": rng.choice(["muy_alto", "alto", "moderado", "bajo"]),
            "alcohol": rng.choice(["frecuente", "ocasional", "raro", "nunca", "nunca"]),
            "tipo_dieta": override.get("tipo_dieta") or rng.choice(["omnivoro"] * 6 + ["pescetariano", "vegetariano", "vegano"]),
            # La app envia cada sintoma marcado como "frecuente" y el resto como "nunca" (buildPayload).
            **{
                symptom: "frecuente" if rng.random() < 0.2 else "nunca"
                for symptom in ("dolor_muscular", "dolor_articular", "niebla_mental", "caida_cabello", "piel_seca", "unas_quebradizas", "calambres")
            },
            "objetivos": rng.sample(["energia", "inmunidad", "suenio", "rendimiento", "salud_osea", "estres", "cabello_piel_unas"], k=rng.randint(1, 3)),
            "tabaco": rng.choice(["nunca", "nunca", "ex", "actual"]),
            "condiciones_seguridad": override.get("condiciones_seguridad", []),
            "restricciones": rng.choice([[], [], [], ["alergia_pescado_mariscos"], ["alergia_lacteos"], ["sin_gluten"]]),
        }
        if rng.random() < 0.8:
            payload.update(_food_block(rng, payload["tipo_dieta"]))
        if rng.random() < 0.2:
            payload["toma_suplementos"] = "si"
            payload["suplementos_actuales"] = rng.sample(
                ["vitamina_d", "magnesio", "hierro", "omega_3", "multivitaminico", "vitamina_b12", "proteina"], k=1,
            )
        labs = [lab for code in override.get("_labs", []) if (lab := _lab_value(code, rng))]
        if labs:
            payload["lab_results"] = labs
        payload.update(_optional_context(rng, payload))
        cases.append({"case_id": f"CR{index + 1:03d}", "payload": payload})
    return cases


_DIET = {"omnivoro": "omnivora", "pescetariano": "pescetariana", "vegetariano": "vegetariana", "vegano": "vegana"}


def describe_profile(payload: dict[str, Any]) -> str:
    """Resumen legible del perfil para quien revisa (en vez de nombres de campos)."""
    parts = [
        f"{'Mujer' if payload['sexo'] == 'femenino' else 'Hombre'} de {payload['age_years']} años, "
        f"{payload['weight_value']} kg, {payload['height_value']} cm",
        f"dieta {_DIET.get(payload['tipo_dieta'], payload['tipo_dieta'])} ({payload['dieta'].replace('_', ' ')})",
        (
            "porciones/sem: " + ", ".join(
                f"{label} {payload[field]}" for field, label in zip(FOOD_FIELDS, ["carne roja", "pollo", "pescado", "huevos", "menestras", "lácteos"])
            ) + f"; frutas y verduras {payload['fruit_veg_servings_day']}/día"
            if "fish_servings_week" in payload else "alimentos: sin responder"
        ),
        f"sueño {payload['horas_sueno'].replace('_', '-')}, fatiga {payload['fatiga'].replace('_', ' ')}, estrés {payload['estres'].replace('_', ' ')}",
        f"sol {payload['exposicion_solar'].replace('_', ' ')}, ejercicio {payload['frecuencia_ejercicio'].replace('_', ' ')}",
        f"alcohol {payload['alcohol']}, tabaco {payload.get('tabaco') or 'sin dato'}",
    ]
    symptoms = [name.replace("_", " ") + f" {payload[name]}" for name in ("caida_cabello", "calambres", "niebla_mental") if payload.get(name) not in (None, "nunca")]
    if symptoms:
        parts.append("síntomas: " + ", ".join(symptoms))
    parts.append("objetivos: " + ", ".join(payload.get("objetivos") or ["ninguno"]))
    context = [
        f"{label}: {str(payload[key]).replace('_', ' ')}"
        for key, label in [
            ("fatiga_duracion", "fatiga desde"), ("fatiga_mejora_descanso", "mejora al descansar"),
            ("protector_solar", "protector solar"), ("piel_expuesta", "piel al sol"),
            ("menstruacion_abundante", "menstruación abundante"), ("parto_ultimo_anio", "parto último año"),
            ("dona_sangre", "dona sangre"), ("anios_dieta_vegetal", "años dieta vegetal"),
            ("alimentos_fortificados_b12", "fortificados con B12"),
        ]
        if payload.get(key)
    ]
    if payload.get("medicamentos_b12"):
        context.append("medicamentos: " + ", ".join(payload["medicamentos_b12"]))
    if context:
        parts.append(" · ".join(context))
    if payload.get("condiciones_seguridad"):
        parts.append("SEGURIDAD: " + ", ".join(item.replace("_", " ") for item in payload["condiciones_seguridad"]))
    if payload.get("restricciones"):
        parts.append("restricciones: " + ", ".join(item.replace("_", " ") for item in payload["restricciones"]))
    if payload.get("suplementos_actuales"):
        parts.append("ya toma: " + ", ".join(item.replace("_", " ") for item in payload["suplementos_actuales"]))
    if payload.get("lab_results"):
        parts.append("LABS: " + "; ".join(f"{lab['code']} {lab['value']} {lab['unit']}" for lab in payload["lab_results"]))
    return " · ".join(parts)


def recommendation_rows(case_id: str, response: dict[str, Any]) -> list[dict[str, Any]]:
    """Formato largo: una fila por componente recomendado (la hoja Casos solo muestra los primeros)."""
    rows = []
    for rank, rec in enumerate(response.get("recommendations") or [], start=1):
        rows.append({
            "case_id": case_id,
            "rank": rank,
            "componente": rec.get("display_name") or rec.get("name"),
            "component_id": rec.get("component_id"),
            "condicion": rec.get("condition_display") or rec.get("condition"),
            "rol": rec.get("recommendation_role") or rec.get("priority"),
            "evidencia": rec.get("evidence_grade") or rec.get("evidence_strength"),
            "compra_permitida": "si" if rec.get("commercial_eligible") else "no",
            "motivo_sin_compra": rec.get("commercial_block_reason") or "",
            "dosis_orientativa": rec.get("dosage_hint") or "",
            "ya_toma": "si" if rec.get("already_taking") else "",
        })
    return rows


def summarize_output(response: dict[str, Any]) -> dict[str, Any]:
    conditions = [
        f"{item.get('display_name') or item.get('code')} ({float(item.get('risk_index', item.get('probability')) or 0):.2f})"
        for item in response.get("conditions_display") or []
    ]
    labs = [
        f"{item.get('label')}: {float(item.get('probability') or 0):.2f}{' (+)' if item.get('positive') else ''}"
        for item in response.get("lab_risk_estimates") or []
    ]
    recs = []
    all_recs = response.get("recommendations") or []
    for index, rec in enumerate(all_recs[:TOP_RECS_IN_CASE], start=1):
        commercial = "compra permitida" if rec.get("commercial_eligible") else f"sin compra: {rec.get('commercial_block_reason') or 'bloqueado'}"
        evidence = rec.get("evidence_grade") or rec.get("evidence_strength") or "s/d"
        recs.append(f"{index}. {rec.get('display_name') or rec.get('name')} — {rec.get('recommendation_role') or rec.get('priority')} — evidencia {evidence} — {commercial}")
    if len(all_recs) > TOP_RECS_IN_CASE:
        recs.append(f"... y {len(all_recs) - TOP_RECS_IN_CASE} más (hoja 'Recomendaciones')")
    alerts = list(response.get("profile_warnings") or []) + list(response.get("safety_actions") or [])
    alerts += [
        f"Interacción: {item.get('component_a')} + {item.get('component_b')}" for item in response.get("alertas") or []
        if isinstance(item, dict)
    ]
    return {
        "sistema_condiciones": "\n".join(conditions) or "(ninguna)",
        "sistema_estimaciones_nhanes": "\n".join(labs),
        "sistema_nivel_seguridad": response.get("safety_level") or "",
        "sistema_compra_bloqueada": "si" if response.get("commercial_recommendations_blocked") else "no",
        "sistema_alertas": "\n".join(alerts),
        "sistema_n_condiciones": len(conditions),
        "sistema_n_recomendaciones": len(all_recs),
        "sistema_recomendaciones": "\n".join(recs) or "(ninguna)",
    }


def run_cases(
    cases: list[dict[str, Any]], recommend: Callable[[dict[str, Any]], dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    rows: list[dict[str, Any]] = []
    long_rows: list[dict[str, Any]] = []
    for case in cases:
        try:
            response = recommend(case["payload"])
            output = summarize_output(response)
            long_rows += recommendation_rows(case["case_id"], response)
            error = ""
        except Exception as exc:  # un caso roto no debe tumbar el paquete; queda visible para el equipo
            output = {}
            error = f"{type(exc).__name__}: {exc}"
        rows.append({
            "case_id": case["case_id"],
            "perfil": describe_profile(case["payload"]),
            **output,
            "error_generacion": error,
            **{column: "" for column in CASE_REVIEW_COLUMNS},
        })
    return rows, long_rows


def default_recommender() -> Callable[[dict[str, Any]], dict[str, Any]]:
    from app.domains.recommendations.servicio_recomendaciones import RecommendationService
    from app.domains.survey.esquemas_encuesta import EncuestaInput
    from app.ml.model_loader import load_all_models

    models = load_all_models()
    if models.get("pipeline_vitaminas") is None:
        raise RuntimeError("Modelos no cargados: faltan artefactos en MODEL_DIR (ver /health).")
    service = RecommendationService(models=models)
    return lambda payload: service.recommend(EncuestaInput(**payload))


# ---------------------------------------------------------------------------
# Reglas
# ---------------------------------------------------------------------------

def rule_frames(knowledge_dir: Path = KNOWLEDGE_DIR) -> dict[str, pd.DataFrame]:
    frames: dict[str, pd.DataFrame] = {}
    for sheet, csv_name, key_cols in RULE_SHEETS:
        df = pd.read_csv(knowledge_dir / csv_name, dtype=str, keep_default_na=False)
        df.insert(0, "rule_key", [row_key(sheet, row, key_cols) for row in df.to_dict("records")])
        for column in RULE_REVIEW_COLUMNS:
            df[column] = ""
        frames[sheet] = df
    from app.domains.catalog.seguridad_productos import DEFAULT_INGREDIENT_SAFETY_RULES

    ingredient = pd.DataFrame([dataclasses.asdict(rule) for rule in DEFAULT_INGREDIENT_SAFETY_RULES]).astype(str)
    ingredient.insert(0, "rule_key", [row_key(INGREDIENT_SHEET, row, ["name"]) for row in ingredient.to_dict("records")])
    for column in RULE_REVIEW_COLUMNS:
        ingredient[column] = ""
    frames[INGREDIENT_SHEET] = ingredient
    return frames


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

def write_workbook(
    path: Path,
    case_rows: list[dict[str, Any]],
    rules: dict[str, pd.DataFrame],
    meta: dict[str, Any],
    recommendation_long: list[dict[str, Any]] | None = None,
) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    review_fill = PatternFill("solid", fgColor="FFF4CC")
    header_font = Font(bold=True)
    path.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        pd.DataFrame({"Instrucciones": INSTRUCTIONS}).to_excel(writer, sheet_name="Instrucciones", index=False)
        pd.DataFrame(case_rows).to_excel(writer, sheet_name="Casos", index=False)
        long_df = pd.DataFrame(recommendation_long) if recommendation_long else pd.DataFrame(columns=["case_id", "rank", "componente"])
        long_df.to_excel(writer, sheet_name="Recomendaciones", index=False)
        for sheet, df in rules.items():
            df.to_excel(writer, sheet_name=sheet, index=False)
        signature = pd.DataFrame({"Campo": SIGNATURE_FIELDS, "Valor": [""] * len(SIGNATURE_FIELDS)})
        signature.to_excel(writer, sheet_name="Firma", index=False)
        manifest = pd.DataFrame(meta["files"])
        manifest.to_excel(writer, sheet_name="Firma", index=False, startrow=len(SIGNATURE_FIELDS) + 3)
        pd.DataFrame([{"key": key, "value": json.dumps(value, ensure_ascii=False)} for key, value in meta.items()]).to_excel(
            writer, sheet_name="_meta", index=False,
        )

        book = writer.book
        book["Instrucciones"].column_dimensions["A"].width = 140
        for cell in book["Instrucciones"]["A"]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        book["_meta"].sheet_state = "hidden"
        book["Firma"].column_dimensions["A"].width = 48
        book["Firma"].column_dimensions["B"].width = 40

        sheets = {"Casos": CASE_REVIEW_COLUMNS, **{sheet: RULE_REVIEW_COLUMNS for sheet in rules}}
        for sheet_name, review_columns in sheets.items():
            ws = book[sheet_name]
            headers = [cell.value for cell in ws[1]]
            ws.freeze_panes = "B2"
            for col_index, header in enumerate(headers, start=1):
                letter = get_column_letter(col_index)
                ws.cell(row=1, column=col_index).font = header_font
                is_review = header in review_columns
                long_text = header in {"perfil", "sistema_recomendaciones", "sistema_condiciones", "sistema_alertas", "message", "rationale", "guidance_note"}
                ws.column_dimensions[letter].width = 60 if long_text else (22 if is_review else 16)
                if long_text:
                    for cell in ws[letter][1:]:
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
                if not is_review:
                    continue
                for cell in ws[letter][1:]:
                    cell.fill = review_fill
                options = review_columns[header]
                if options and ws.max_row > 1:
                    validation = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
                    validation.add(f"{letter}2:{letter}{ws.max_row}")
                    ws.add_data_validation(validation)


def build_meta(n_cases: int, seed: int) -> dict[str, Any]:
    files = file_manifest()
    fingerprint = hashlib.sha256(json.dumps(files, sort_keys=True).encode()).hexdigest()[:12]
    return {
        "package_version": PACKAGE_VERSION,
        "package_id": f"{PACKAGE_VERSION}_{datetime.now(timezone.utc):%Y%m%d}_{fingerprint}",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "cases": n_cases,
        "seed": seed,
        "files": files,
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Genera el paquete de revision clinica (.xlsx + manifest).")
    parser.add_argument("--cases", type=int, default=120, help="Numero de casos (plan: >=100).")
    parser.add_argument("--seed", type=int, default=2026)
    parser.add_argument("--out-dir", type=Path, default=DEFAULT_OUT_DIR)
    return parser


def main(argv: list[str] | None = None, recommend: Callable[[dict[str, Any]], dict[str, Any]] | None = None) -> int:
    args = build_parser().parse_args(argv)
    if args.cases < 100:
        print(f"AVISO: {args.cases} casos; el plan pide >=100 para dar la revision por valida.", file=sys.stderr)
    cases = build_case_payloads(args.cases, seed=args.seed)
    case_rows, recommendation_long = run_cases(cases, recommend or default_recommender())
    rules = rule_frames()
    meta = build_meta(len(case_rows), args.seed)

    xlsx_path = args.out_dir / f"{meta['package_id']}.xlsx"
    write_workbook(xlsx_path, case_rows, rules, meta, recommendation_long)
    manifest_path = xlsx_path.with_suffix(".manifest.json")
    manifest_path.write_text(json.dumps({**meta, "rule_rows": {sheet: len(df) for sheet, df in rules.items()}}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    failed = [row["case_id"] for row in case_rows if row["error_generacion"]]
    print(json.dumps({
        "xlsx": str(xlsx_path),
        "manifest": str(manifest_path),
        "cases": len(case_rows),
        "case_errors": failed,
        "rule_rows": {sheet: len(df) for sheet, df in rules.items()},
    }, ensure_ascii=False, indent=2))
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
